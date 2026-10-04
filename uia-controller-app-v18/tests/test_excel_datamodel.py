"""
Test Suite for UiA Enterprise Excel Data Model Generator
Verifies sheet structure, table registrations, formula syntax, and data integrity.
"""

import os
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src" / "tools"))

import openpyxl
import pytest
from generate_excel_datamodel import create_excel_datamodel, EXCEL_FILE_PRIMARY, EXCEL_FILE_SECONDARY

def test_excel_datamodel_generation():
    # 1. Run generation
    filepath = create_excel_datamodel()
    assert os.path.exists(filepath), "Primary Excel Data Model file was not created"
    assert os.path.exists(EXCEL_FILE_SECONDARY), "Secondary report mirror file was not created"
    
    # 2. Verify Workbook and Sheets
    wb = openpyxl.load_workbook(filepath, data_only=False)
    expected_sheets = [
        "00_Navigasjon_Arkitektur",
        "01_Executive_Cockpit",
        "02_EVM_Prosjektkontroll",
        "03_S_Kurve_Tidsserie",
        "04_F05_20_Avsetningskalkyle",
        "05_Compliance_Internkontroll",
        "Dim_Project",
        "Dim_Date",
        "Dim_Model_Parameter",
        "Fact_EVM_Snapshots",
        "Fact_UBW_Audit",
        "Fact_Travel_Audit",
        "Fact_Project_BOA",
        "Dim_Kontoplan_2026"
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Missing sheet: {s}"
        
    # 3. Verify Table Registrations (Star Schema Tables)
    tables_map = {
        "Dim_Project": "tbl_Dim_Project",
        "Dim_Date": "tbl_Dim_Date",
        "Dim_Model_Parameter": "tbl_Dim_Parameters",
        "Fact_EVM_Snapshots": "tbl_Fact_EVM",
        "Fact_UBW_Audit": "tbl_Fact_UBW",
        "Fact_Travel_Audit": "tbl_Fact_Travel",
        "Fact_Project_BOA": "tbl_Fact_BOA",
        "Dim_Kontoplan_2026": "tbl_Dim_Kontoplan"
    }
    for sheet_name, tab_name in tables_map.items():
        ws = wb[sheet_name]
        registered_tables = [t.name for t in ws.tables.values()]
        assert tab_name in registered_tables, f"Table {tab_name} not registered in {sheet_name}"
        
    # 4. Verify EVM formulas in 02_EVM_Prosjektkontroll
    ws_evm = wb["02_EVM_Prosjektkontroll"]
    assert ws_evm["G5"].value == "=E5-F5", "CV formula mismatch"
    assert ws_evm["H5"].value == "=E5-D5", "SV formula mismatch"
    assert "=IFERROR(E5/F5, 1.00)" in ws_evm["I5"].value, "CPI formula mismatch"
    assert "=IFERROR(E5/D5, 1.00)" in ws_evm["J5"].value, "SPI formula mismatch"
    
    # 5. Verify F-05-20 formulas in 04_F05_20_Avsetningskalkyle
    ws_f05 = wb["04_F05_20_Avsetningskalkyle"]
    assert ws_f05["B7"].value == "=B6*0.05", "5% cap formula mismatch"
    assert ws_f05["B9"].value == "=MAX(0, B8-B7)", "Excess carryover formula mismatch"
    
    # 6. Verify Chart in 03_S_Kurve_Tidsserie
    ws_ts = wb["03_S_Kurve_Tidsserie"]
    assert len(ws_ts._charts) > 0, "Line chart missing in S-curve sheet"
